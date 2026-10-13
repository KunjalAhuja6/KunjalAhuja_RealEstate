"""
Presentation Generator (PPTX & PDF)
Generates:
1. presentation/Project_Presentation.pptx (10 slides max via python-pptx)
2. presentation/Project_Presentation.pdf (matching 10-slide deck via ReportLab)
Author / Presenter: Kunjal Ahuja
Prodapt Capstone Project
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas


# ==========================================
# SLIDE DATA DEFINITION (10 SLIDES EXACTLY)
# ==========================================
SLIDES_DATA = [
    {
        "num": 1,
        "title": "AI-Powered Real Estate Property Valuation Assistant",
        "subtitle": "Prodapt Capstone Project • Panel Presentation (8 Min Demo + 2 Min Q&A)",
        "bullets": [
            "Project Domain: Real Estate & PropTech Valuation Intelligence",
            "Author & Presenter: Kunjal Ahuja (Senior AI/ML Engineer)",
            "Architecture: Autonomous Multi-Agent Collaborative System with Hybrid RAG & Supervised ML",
            "Core Technologies: FastAPI, React 18, Scikit-learn (Gradient Boosting R² = 0.9198), SentenceTransformers, DeepEval (v4.2.8)"
        ],
        "metrics": [
            ("10,000", "Dataset Records"),
            ("0.9198", "Model Test R²"),
            ("100%", "DeepEval Pass Rate"),
            ("16", "REST Endpoints")
        ]
    },
    {
        "num": 2,
        "title": "Problem Statement & Business Need",
        "subtitle": "Overcoming Inefficiencies and Subjectivity in Traditional Real Estate Appraisals",
        "bullets": [
            "Manual Search Inefficiency: Traditional appraisals require scanning thousands of historical MLS records, taking 3–7 business days.",
            "Subjective Valuation Adjustments: Ad-hoc manual adjustments for condition and age introduce severe appraiser bias.",
            "Unquantified Physical & Market Risk: Properties with deferred maintenance lack standardized financial risk quantification.",
            "Business Objective: Build an autonomous valuation platform delivering instant comp retrieval, ML predictions, risk scoring, and LLM-as-judge audit."
        ],
        "metrics": [
            ("3-7 Days", "Manual Appraisal Latency"),
            ("< 500 ms", "System Response Time"),
            ("0-100", "Quantified Risk Index"),
            ("14 Fields", "Automated Guardrails")
        ]
    },
    {
        "num": 3,
        "title": "Solution Overview & High-Level Architecture",
        "subtitle": "End-to-End Intelligence Pipeline from Ground Truth to Appraiser Dossier",
        "bullets": [
            "1. Data Pipeline: Standardized 10,000-record Ames dataset with 15 key physical fields and 1-10 quality scale.",
            "2. True Hybrid RAG: 384-d dense vectors (all-MiniLM-L6-v2) + full candidate set TF-IDF keyword scoring (alpha = 0.55).",
            "3. Supervised ML: GradientBoostingRegressor with 90% empirical prediction intervals (Test R² = 0.9198, MAPE = 8.87%).",
            "4. Multi-Agent Collaboration: Specialized agents for classification, valuation, risk scoring, and LLM-as-judge auditing.",
            "5. Production Delivery: FastAPI REST microservice (16 endpoints) + modern React 18 single-page application (7 tabs)."
        ],
        "metrics": [
            ("Hybrid RAG", "Dense + TF-IDF"),
            ("GradientBoosting", "R² = 0.9198"),
            ("React 18", "7 Workflow Tabs"),
            ("FastAPI", "16 Endpoints")
        ]
    },
    {
        "num": 4,
        "title": "Dataset & Gaussian Copula Synthetic Augmentation",
        "subtitle": "Honest Data Engineering with Strict Real-Only Test Evaluation",
        "bullets": [
            "Historical Foundation: 2,930 real residential sales from Ames Housing Dataset (OpenML 43926 / Dean De Cock 2011).",
            "Gaussian Copula Augmentation: 7,070 synthetic records fitted per neighborhood on empirical marginal distributions.",
            "Novelty Distance Filter: Rejects synthetic points closer than 25th percentile of real nearest neighbors, preventing memorization.",
            "Honest Real-Only Test Set: 20% of real sales (n = 586, random_state=42). Synthetic records never appear in the test set.",
            "Comparable Isolation: Comps searches strictly exclude synthetic records by default (include_synthetic=False)."
        ],
        "metrics": [
            ("2,930", "Real Historical Sales"),
            ("7,070", "Synthetic Records"),
            ("586", "Real-Only Test Set"),
            ("28", "Ames Neighborhoods")
        ]
    },
    {
        "num": 5,
        "title": "Hybrid RAG & Dynamic Comparable Prioritization",
        "subtitle": "Grounded Evidence Retrieval with Automated Neighborhood Widening",
        "bullets": [
            "Dense Semantic Retrieval: In-memory normalized NumPy cosine index over cached 384-d embeddings (sub-5ms query).",
            "Lexical TF-IDF Scoring: Evaluates full metadata-filtered candidate set before fusion, preventing silent drops.",
            "Multi-Factor Comparability Index: Location (30%), Living Area (25%), Quality (20%), Age (15%), Rooms (10%).",
            "Dynamic Neighborhood Widening: Automatically widens search citywide when < 3 comps in neighborhood match living area +/- 25% and quality +/- 1.",
            "Strict Zero-Match Handling: Returns empty list [] on zero matches rather than falling back to unconstrained properties."
        ],
        "metrics": [
            ("alpha = 0.55", "Dense / Lexical Split"),
            ("5 Factors", "Comparability Index"),
            ("Citywide", "Dynamic Widening"),
            ("< 5 ms", "In-Memory Cosine")
        ]
    },
    {
        "num": 6,
        "title": "Supervised ML Modeling & What-If Simulation",
        "subtitle": "High-Precision Valuation Benchmark with Direct Feature Pass-Through",
        "bullets": [
            "Active Regressor: GradientBoostingRegressor achieved Test R² = 0.9198, RMSE = $25,361.84, MAE = $16,161.76, MAPE = 8.87%.",
            "Baseline Comparison: Outperformed real-only baseline model (R² = 0.9031, RMSE = $27,870.39), demonstrating augmentation benefit.",
            "90% Prediction Intervals: Derived empirically from test residuals on real unseen transactions.",
            "What-If Renovation Simulator: Simulates value lift from added square footage, room additions, and finish upgrades in real time.",
            "Top Valuation Drivers: Overall Quality (54%), Above-Grade Living Area (22%), Total Basement Area (8%), Location Premiums."
        ],
        "metrics": [
            ("0.9198", "Test R² Score"),
            ("$25,361", "Test RMSE"),
            ("8.87%", "Test MAPE"),
            ("90%", "Prediction Interval")
        ]
    },
    {
        "num": 7,
        "title": "Autonomous Multi-Agent Workflow & LLM-as-Judge",
        "subtitle": "Collaborative Reasoning Core with Dynamic comp_fit Weighting",
        "bullets": [
            "Guardrail Agent: Enforces 14 physical boundary constraints (100-15k sqft, 1-10 qual) and detects prompt injection.",
            "Classification Agent: Infers Market Tier, Typology, Condition Grade, and Target Demographics.",
            "Basic Valuation Agent: Computes median comparable price ($) and size-weighted PPSF valuation.",
            "Dynamic Reconciliation: comp_fit weighting (comp_weight: 0.1–0.5, ml_weight: 0.5–0.9) replaces static 50/50 splits.",
            "Property Risk Agent: 0-100 Risk Index (Physical Age, Condition Gap, Neighborhood Volatility, Layout Liquidity).",
            "LLM-as-Judge Auditor: Audits valuation against comps and ML bounds, issuing formal verdicts (APPROVED)."
        ],
        "metrics": [
            ("comp_fit", "Dynamic Weighting"),
            ("APPROVED", "Judge Audit Verdict"),
            ("4 Factors", "Risk Index Breakdown"),
            ("Dual-Mode", "Live LLM + Fallback")
        ]
    },
    {
        "num": 8,
        "title": "Evaluation Results (DeepEval Framework)",
        "subtitle": "Rigorous Recommendation Quality Verification across 5 Market Scenarios",
        "bullets": [
            "Framework: DeepEval (v4.2.8 pinned) executed in live_llm_judge mode across 5 curated appraisal scenarios.",
            "Benchmark Pass Rate: 100.0% (5 of 5 Passed) with total execution time of 61.98 seconds.",
            "Faithfulness (0.900): High factual grounding; valuation narratives cite verified historical comps without hallucination.",
            "Answer Relevancy (0.880): Responses directly address user query parameters and physical constraints.",
            "Contextual Precision & Recall (0.850 / 0.850): Optimal comparable ranking and comprehensive coverage of valuation dimensions.",
            "Valuation Accuracy (1.000): Reconciled values fell strictly within verified historical ground-truth bounds in all scenarios."
        ],
        "metrics": [
            ("100%", "Pass Rate (5/5)"),
            ("0.900", "Faithfulness Score"),
            ("0.880", "Answer Relevancy"),
            ("1.000", "Valuation Accuracy")
        ]
    },
    {
        "num": 9,
        "title": "Continuous Learning Feedback Loop & Calibration",
        "subtitle": "Automated Dynamic Calibration of Systemic Pricing Bias",
        "bullets": [
            "Appraiser Feedback Logging: Certified appraisers submit verified closing prices, 1-5 star ratings, and qualitative notes.",
            "Persistent Storage: Feedback records stored in data/sample_data/feedback.json with error tracking.",
            "Uncalibrated Bias Tracking: Errors and systematic bias are evaluated against uncalibrated estimates (recommended / applied_factor).",
            "Drift Prevention: Prevents successive calibration factors from cancelling out and drifting back to 1.0.",
            "Real-Time Multiplier: Dynamically adjusts future reconciled valuations via active calibration factor."
        ],
        "metrics": [
            ("feedback.json", "Persistent Store"),
            ("Uncalibrated", "Bias Tracking"),
            ("Real-Time", "Dynamic Calibration"),
            ("0 Drift", "Stable Convergence")
        ]
    },
    {
        "num": 10,
        "title": "Production Deployment, Testing & Panel Q&A",
        "subtitle": "Complete Deliverables, Automated Verification & Discussion",
        "bullets": [
            "Production React 18 UI: 7 accessible workflow tabs (Valuation, Comps, ML, What-If, Risk & Audit, Feedback, Evaluation).",
            "FastAPI REST Microservice: 16 documented endpoints with interactive OpenAPI Swagger docs at /docs.",
            "Automated Test Suite: 36 pytest unit and integration tests passing cleanly (100% test pass rate).",
            "Deliverables Complete: Architecture Diagram (PDF/PNG), Design Document (PDF), API Docs (PDF), Dataset Details (PDF), Evaluation Report (PDF), Slide Deck (PPTX/PDF).",
            "Open for Panel Discussion: Ready for live demonstration and panel Q&A. Thank you!"
        ],
        "metrics": [
            ("36 Tests", "100% Passing"),
            ("16 Endpoints", "FastAPI REST"),
            ("7 Tabs", "React 18 Frontend"),
            ("Kunjal Ahuja", "Presenter / Author")
        ]
    }
]


# ==========================================
# 1. GENERATE PPTX FILE (python-pptx)
# ==========================================
def generate_pptx(pptx_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide

    c_navy = RGBColor(30, 58, 138)
    c_teal = RGBColor(15, 118, 110)
    c_slate = RGBColor(30, 41, 59)
    c_bg = RGBColor(248, 250, 252)
    c_white = RGBColor(255, 255, 255)
    c_gray = RGBColor(100, 116, 139)
    c_card_border = RGBColor(203, 213, 225)

    for slide_data in SLIDES_DATA:
        slide = prs.slides.add_slide(blank_layout)
        
        # Background rect
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_bg
        bg.line.color.rgb = c_bg

        # Top Banner
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2))
        banner.fill.solid()
        banner.fill.fore_color.rgb = c_navy
        banner.line.color.rgb = c_navy

        # Title text
        tx_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.12), Inches(12.133), Inches(0.55))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = slide_data["title"]
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = c_white

        # Subtitle text
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.65), Inches(12.133), Inches(0.45))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = slide_data["subtitle"]
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = RGBColor(191, 219, 254)

        # Main Content Card (Left)
        card_main = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.5), Inches(8.5), Inches(5.4))
        card_main.fill.solid()
        card_main.fill.fore_color.rgb = c_white
        card_main.line.color.rgb = c_card_border
        card_main.line.width = Pt(1.5)

        tx_main = slide.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(8.1), Inches(5.1))
        tf_main = tx_main.text_frame
        tf_main.word_wrap = True

        for b_idx, bullet in enumerate(slide_data["bullets"]):
            p_bullet = tf_main.add_paragraph() if b_idx > 0 else tf_main.paragraphs[0]
            p_bullet.text = f"•  {bullet}"
            p_bullet.font.size = Pt(12)
            p_bullet.font.color.rgb = c_slate
            p_bullet.space_after = Pt(12)

        # Metrics Column (Right) - 4 Metric Cards
        metrics = slide_data.get("metrics", [])
        m_y = 1.5
        m_h = 1.18
        for m_val, m_lbl in metrics:
            m_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.3), Inches(m_y), Inches(3.433), Inches(m_h))
            m_card.fill.solid()
            m_card.fill.fore_color.rgb = c_white
            m_card.line.color.rgb = c_teal
            m_card.line.width = Pt(1.5)

            tx_m = slide.shapes.add_textbox(Inches(9.4), Inches(m_y + 0.1), Inches(3.233), Inches(m_h - 0.2))
            tf_m = tx_m.text_frame
            tf_m.word_wrap = True
            
            p_val = tf_m.paragraphs[0]
            p_val.text = m_val
            p_val.font.size = Pt(20)
            p_val.font.bold = True
            p_val.font.color.rgb = c_teal
            p_val.alignment = PP_ALIGN.CENTER

            p_lbl = tf_m.add_paragraph()
            p_lbl.text = m_lbl
            p_lbl.font.size = Pt(10)
            p_lbl.font.color.rgb = c_slate
            p_lbl.alignment = PP_ALIGN.CENTER

            m_y += m_h + 0.22

        # Footer line & text
        ft_box = slide.shapes.add_textbox(Inches(0.6), Inches(7.05), Inches(12.133), Inches(0.35))
        tf_ft = ft_box.text_frame
        p_ft = tf_ft.paragraphs[0]
        p_ft.text = f"AI-Powered Real Estate Property Valuation Assistant • Slide {slide_data['num']} of 10 • Presenter: Kunjal Ahuja"
        p_ft.font.size = Pt(9)
        p_ft.font.color.rgb = c_gray

    prs.save(pptx_path)
    print(f"Successfully generated {pptx_path} (size: {os.path.getsize(pptx_path)} bytes)")


# ==========================================
# 2. GENERATE PDF FILE (ReportLab)
# ==========================================
def generate_pdf(pdf_path):
    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=landscape(letter),
        leftMargin=36, rightMargin=36,
        topMargin=28, bottomMargin=28
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'SlideTitle', parent=styles['Heading1'],
        fontName='Helvetica-Bold', fontSize=18, leading=22,
        textColor=colors.HexColor('#1e3a8a')
    )
    subtitle_style = ParagraphStyle(
        'SlideSub', parent=styles['Normal'],
        fontName='Helvetica', fontSize=10, leading=14,
        textColor=colors.HexColor('#0f766e'), spaceAfter=8
    )
    bullet_style = ParagraphStyle(
        'SlideBullet', parent=styles['Normal'],
        fontName='Helvetica', fontSize=8.5, leading=12.5,
        textColor=colors.HexColor('#1e293b'), spaceAfter=5
    )
    metric_val_style = ParagraphStyle(
        'MetricVal', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=14, leading=17,
        textColor=colors.HexColor('#0f766e'), alignment=1
    )
    metric_lbl_style = ParagraphStyle(
        'MetricLbl', parent=styles['Normal'],
        fontName='Helvetica', fontSize=7.5, leading=10,
        textColor=colors.HexColor('#475569'), alignment=1
    )
    footer_style = ParagraphStyle(
        'SlideFooter', parent=styles['Normal'],
        fontName='Helvetica', fontSize=7.5, leading=9,
        textColor=colors.HexColor('#64748b')
    )

    story = []

    for idx, slide_data in enumerate(SLIDES_DATA):
        if idx > 0:
            story.append(PageBreak())

        # Title & Subtitle
        story.append(Paragraph(slide_data["title"], title_style))
        story.append(Paragraph(slide_data["subtitle"], subtitle_style))

        # Main bullet points (Left) & Metrics (Right)
        bullet_flowables = []
        for bullet in slide_data["bullets"]:
            bullet_flowables.append(Paragraph(f"•  {bullet}", bullet_style))

        # Metric boxes (Right column)
        metric_boxes = []
        for m_val, m_lbl in slide_data["metrics"]:
            mv = Paragraph(m_val, metric_val_style)
            ml = Paragraph(m_lbl, metric_lbl_style)
            m_t = Table([[mv], [ml]], colWidths=[170])
            m_t.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
                ('BOX', (0,0), (-1,-1), 0.8, colors.HexColor('#0f766e')),
                ('TOPPADDING', (0,0), (-1,-1), 4),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ]))
            metric_boxes.append(m_t)
            metric_boxes.append(Spacer(1, 4))

        # 2-Column Table Layout: Left (bullets, 520 pt) | Right (metrics, 180 pt)
        col_left = Table([[b] for b in bullet_flowables], colWidths=[520])
        col_left.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.white),
            ('BOX', (0,0), (-1,-1), 0.8, colors.HexColor('#cbd5e1')),
            ('TOPPADDING', (0,0), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 8),
            ('LEFTPADDING', (0,0), (-1,-1), 10),
            ('RIGHTPADDING', (0,0), (-1,-1), 10),
        ]))

        col_right = Table([[mb] for mb in metric_boxes], colWidths=[180])
        col_right.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('LEFTPADDING', (0,0), (-1,-1), 4),
            ('RIGHTPADDING', (0,0), (-1,-1), 4),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ]))

        slide_grid = Table([[col_left, col_right]], colWidths=[530, 190])
        slide_grid.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('TOPPADDING', (0,0), (-1,-1), 0),
            ('BOTTOMPADDING', (0,0), (-1,-1), 0),
            ('LEFTPADDING', (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
        ]))

        story.append(slide_grid)
        story.append(Spacer(1, 10))

        # Footer
        ft_text = f"AI-Powered Real Estate Property Valuation Assistant • Slide {slide_data['num']} of 10 • Presenter: Kunjal Ahuja • Prodapt Capstone"
        story.append(Paragraph(ft_text, footer_style))

    doc.build(story)
    print(f"Successfully generated {pdf_path} (size: {os.path.getsize(pdf_path)} bytes)")


if __name__ == '__main__':
    pptx = 'presentation/Project_Presentation.pptx'
    pdf = 'presentation/Project_Presentation.pdf'
    generate_pptx(pptx)
    generate_pdf(pdf)
